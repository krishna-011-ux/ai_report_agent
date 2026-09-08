"""
Excel Report Generator
-----------------------
Cleaned data + AI insights ko ek professional formatted .xlsx file mein
convert karta hai, jisme:
  - Ek sheet raw/cleaned data ke liye
  - Ek sheet AI insights/summary ke liye
  - Auto column-width, header styling, aur ek native Excel chart
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, Reference
from datetime import datetime


def generate_excel_report(df: pd.DataFrame, insights: dict, report_config: dict) -> str:
    wb = Workbook()

    # ---------- Sheet 1: Summary & Insights ----------
    ws_summary = wb.active
    ws_summary.title = "Summary"

    title = report_config.get("title", "Automated Report")
    ws_summary["A1"] = title
    ws_summary["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws_summary["A1"].fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    ws_summary.merge_cells("A1:D1")

    ws_summary["A2"] = f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws_summary["A2"].font = Font(italic=True, size=9)

    row = 4
    ws_summary[f"A{row}"] = "Overall Summary"
    ws_summary[f"A{row}"].font = Font(bold=True, size=12)
    row += 1
    ws_summary[f"A{row}"] = insights.get("summary", "")
    ws_summary.merge_cells(f"A{row}:D{row}")
    ws_summary[f"A{row}"].alignment = Alignment(wrap_text=True)
    row += 2

    ws_summary[f"A{row}"] = "Key Insights"
    ws_summary[f"A{row}"].font = Font(bold=True, size=12)
    row += 1
    for insight in insights.get("insights", []):
        ws_summary[f"A{row}"] = f"• {insight}"
        ws_summary.merge_cells(f"A{row}:D{row}")
        ws_summary[f"A{row}"].alignment = Alignment(wrap_text=True)
        row += 1

    row += 1
    ws_summary[f"A{row}"] = "Recommendations"
    ws_summary[f"A{row}"].font = Font(bold=True, size=12)
    row += 1
    for rec in insights.get("recommendations", []):
        ws_summary[f"A{row}"] = f"• {rec}"
        ws_summary.merge_cells(f"A{row}:D{row}")
        ws_summary[f"A{row}"].alignment = Alignment(wrap_text=True)
        row += 1

    ws_summary.column_dimensions["A"].width = 100

    # ---------- Sheet 2: Raw Data ----------
    ws_data = wb.create_sheet("Data")
    for r_idx, row_data in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row_data, 1):
            cell = ws_data.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")

    # auto column width
    for col_cells in ws_data.columns:
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws_data.column_dimensions[col_cells[0].column_letter].width = min(max_len + 3, 30)

    # ---------- Optional native chart (agar numeric + categorical column mile) ----------
    numeric_cols = df.select_dtypes(include="number").columns.tolist()
    text_cols = df.select_dtypes(include="object").columns.tolist()

    if numeric_cols and text_cols:
        try:
            grouped = df.groupby(text_cols[0])[numeric_cols[0]].sum().reset_index()
            ws_chart_data = wb.create_sheet("ChartData")
            ws_chart_data.append([text_cols[0], numeric_cols[0]])
            for _, r in grouped.iterrows():
                ws_chart_data.append([r[text_cols[0]], r[numeric_cols[0]]])

            chart = BarChart()
            chart.title = f"{numeric_cols[0]} by {text_cols[0]}"
            chart.x_axis.title = text_cols[0]
            chart.y_axis.title = numeric_cols[0]

            data_ref = Reference(ws_chart_data, min_col=2, min_row=1, max_row=len(grouped) + 1)
            cats_ref = Reference(ws_chart_data, min_col=1, min_row=2, max_row=len(grouped) + 1)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            ws_data.add_chart(chart, "H2")
        except Exception as e:
            print(f"[ExcelReport] Chart generation skip ho gaya: {e}")

    # ---------- Save ----------
    output_dir = report_config.get("output_dir", "outputs")
    filename = f"{output_dir}/report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    print(f"[ExcelReport] Saved: {filename}")
    return filename
